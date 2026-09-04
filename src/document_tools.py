import re
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader

WORKSPACE = (Path(__file__).parent / "workspace").resolve()

SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".txt", ".md", ".py", ".json", ".csv",
    ".yaml", ".yml", ".html", ".css", ".js", ".ts",
}
TEXT_EXTENSIONS = SUPPORTED_EXTENSIONS - {".pdf", ".docx"}

MAX_READ_CHARS = 40000
MAX_SEARCH_RESULTS = 30

SEARCH_STOP_WORDS = {
    "and", "are", "did", "does", "for", "from", "how", "into", "the",
    "this", "was", "what", "when", "where", "which", "why", "with",
}


def _safe_path(relative_path: str) -> Path:
    if not relative_path or not relative_path.strip():
        raise ValueError("A document path is required.")

    candidate = (WORKSPACE / relative_path).resolve()
    if candidate != WORKSPACE and WORKSPACE not in candidate.parents:
        raise ValueError("That path is outside the allowed document workspace.")
    return candidate


def _relative(path: Path) -> str:
    return str(path.relative_to(WORKSPACE)).replace("\\", "/")


def _supported_files() -> list[Path]:
    if not WORKSPACE.exists():
        return []
    return sorted(
        path for path in WORKSPACE.rglob("*")
        if path.is_file()
        and path.suffix.lower() in SUPPORTED_EXTENSIONS
        and not any(part.startswith(".") for part in path.relative_to(WORKSPACE).parts)
    )


def _pdf_page_text(reader: PdfReader, page_index: int) -> str:
    try:
        return reader.pages[page_index].extract_text() or ""
    except Exception:
        return ""


def list_documents() -> dict[str, Any]:
    """List readable documents in the local workspace."""
    documents = []
    for path in _supported_files():
        item = {
            "path": _relative(path),
            "type": path.suffix.lower().lstrip("."),
            "size_kb": round(path.stat().st_size / 1024, 1),
        }
        if path.suffix.lower() == ".pdf":
            try:
                item["pages"] = len(PdfReader(str(path)).pages)
            except Exception:
                item["pages"] = None
        documents.append(item)

    return {
        "status": "success",
        "workspace": str(WORKSPACE),
        "document_count": len(documents),
        "documents": documents,
    }


def get_document_info(relative_path: str) -> dict[str, Any]:
    """Return metadata about one local document."""
    try:
        path = _safe_path(relative_path)
    except ValueError as exc:
        return {"status": "error", "message": str(exc)}

    if not path.exists() or not path.is_file():
        return {"status": "error", "message": f"Document not found: {relative_path}"}

    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        return {"status": "error", "message": f"Unsupported type: {suffix}"}

    result = {
        "status": "success",
        "path": _relative(path),
        "type": suffix.lstrip("."),
        "size_kb": round(path.stat().st_size / 1024, 1),
    }

    try:
        if suffix == ".pdf":
            reader = PdfReader(str(path))
            result["pages"] = len(reader.pages)
        elif suffix == ".docx":
            result["paragraphs"] = len(Document(str(path)).paragraphs)
        else:
            text = path.read_text(encoding="utf-8", errors="replace")
            result["lines"] = len(text.splitlines())
            result["characters"] = len(text)
    except Exception as exc:
        return {"status": "error", "message": f"Could not inspect file: {exc}"}

    return result


def read_document(relative_path: str) -> dict[str, Any]:
    """Read a local document. Large content is capped to keep tool responses manageable."""
    try:
        path = _safe_path(relative_path)
    except ValueError as exc:
        return {"status": "error", "message": str(exc)}

    if not path.exists() or not path.is_file():
        return {"status": "error", "message": f"Document not found: {relative_path}"}

    suffix = path.suffix.lower()

    try:
        if suffix == ".pdf":
            reader = PdfReader(str(path))
            text = "\n\n".join(
                f"--- Page {i + 1} ---\n{_pdf_page_text(reader, i)}"
                for i in range(len(reader.pages))
            )
        elif suffix == ".docx":
            doc = Document(str(path))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        elif suffix in TEXT_EXTENSIONS:
            text = path.read_text(encoding="utf-8", errors="replace")
        else:
            return {"status": "error", "message": f"Unsupported type: {suffix}"}
    except Exception as exc:
        return {"status": "error", "message": f"Could not read file: {exc}"}

    truncated = len(text) > MAX_READ_CHARS
    return {
        "status": "success",
        "path": _relative(path),
        "content": text[:MAX_READ_CHARS],
        "truncated": truncated,
    }


def read_pdf_pages(
    relative_path: str,
    start_page: int = 1,
    end_page: int = 3,
) -> dict[str, Any]:
    """Read an inclusive PDF page range using 1-based page numbers."""
    try:
        path = _safe_path(relative_path)
    except ValueError as exc:
        return {"status": "error", "message": str(exc)}

    if not path.exists() or not path.is_file():
        return {"status": "error", "message": f"Document not found: {relative_path}"}

    if path.suffix.lower() != ".pdf":
        return {"status": "error", "message": "This tool only supports PDFs."}

    try:
        reader = PdfReader(str(path))
        total = len(reader.pages)
    except Exception as exc:
        return {"status": "error", "message": f"Could not open PDF: {exc}"}

    start_page = max(1, int(start_page))
    end_page = min(max(start_page, int(end_page)), start_page + 7, total)

    if start_page > total:
        return {"status": "error", "message": f"PDF has only {total} pages."}

    content = "\n\n".join(
        f"--- Page {page} ---\n{_pdf_page_text(reader, page - 1)}"
        for page in range(start_page, end_page + 1)
    )

    return {
        "status": "success",
        "path": _relative(path),
        "start_page": start_page,
        "end_page": end_page,
        "total_pages": total,
        "content": content[:MAX_READ_CHARS],
        "truncated": len(content) > MAX_READ_CHARS,
    }


def search_documents(query: str, max_results: int = 12) -> dict[str, Any]:
    """Search all local documents for a word or phrase and return matching snippets."""
    query = (query or "").strip()
    if not query:
        return {"status": "error", "message": "Search query cannot be empty."}

    max_results = max(1, min(int(max_results), MAX_SEARCH_RESULTS))
    needle = query.casefold()
    keywords = [
        token for token in re.findall(r"[\w-]+", needle)
        if len(token) >= 3 and token not in SEARCH_STOP_WORDS
    ]
    matches = []

    def find_match(text: str) -> tuple[int, list[str]] | None:
        folded = text.casefold()
        exact_position = folded.find(needle)
        if exact_position >= 0:
            return exact_position, [query]

        positions = [(folded.find(term), term) for term in keywords]
        hits = [(position, term) for position, term in positions if position >= 0]
        required_hits = 1 if len(keywords) < 3 else 2
        if len(hits) < required_hits:
            return None
        return min(position for position, _ in hits), [term for _, term in hits]

    for path in _supported_files():
        suffix = path.suffix.lower()
        try:
            if suffix == ".pdf":
                reader = PdfReader(str(path))
                for page_num in range(1, len(reader.pages) + 1):
                    text = _pdf_page_text(reader, page_num - 1)
                    found = find_match(text)
                    if found:
                        pos, matched_terms = found
                        start = max(0, pos - 160)
                        end = min(len(text), pos + len(query) + 260)
                        matches.append({
                            "path": _relative(path),
                            "location": f"page {page_num}",
                            "snippet": text[start:end].replace("\n", " ")[:450],
                            "matched_terms": matched_terms,
                        })
                        if len(matches) >= max_results:
                            break

            elif suffix == ".docx":
                doc = Document(str(path))
                for para_num, para in enumerate(doc.paragraphs, start=1):
                    found = find_match(para.text)
                    if found:
                        _, matched_terms = found
                        matches.append({
                            "path": _relative(path),
                            "location": f"paragraph {para_num}",
                            "snippet": para.text[:450],
                            "matched_terms": matched_terms,
                        })
                        if len(matches) >= max_results:
                            break

            elif suffix in TEXT_EXTENSIONS:
                lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
                for line_num, line in enumerate(lines, start=1):
                    found = find_match(line)
                    if found:
                        _, matched_terms = found
                        matches.append({
                            "path": _relative(path),
                            "location": f"line {line_num}",
                            "snippet": line[:450],
                            "matched_terms": matched_terms,
                        })
                        if len(matches) >= max_results:
                            break
        except Exception:
            continue

        if len(matches) >= max_results:
            break

    return {
        "status": "success",
        "query": query,
        "match_count": len(matches),
        "matches": matches,
        "truncated": len(matches) >= max_results,
    }
